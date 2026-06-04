"""Document ingestion helpers for Streamlit uploads."""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any

@dataclass
class LoadedDocument:
    name: str
    text: str


def load_txt_bytes(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def load_docx_bytes(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def load_pdf_bytes(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def load_xlsx_bytes(data: bytes) -> str:
    import pandas as pd
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    parts = []
    for name, frame in sheets.items():
        parts.append(f"Sheet: {name}")
        parts.append(frame.fillna("").astype(str).to_string(index=False, header=False))
    return "\n".join(parts)


def load_uploaded_file(uploaded_file: Any) -> LoadedDocument:
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    data = uploaded_file.read()
    if suffix == ".txt":
        text = load_txt_bytes(data)
    elif suffix == ".docx":
        text = load_docx_bytes(data)
    elif suffix == ".pdf":
        text = load_pdf_bytes(data)
    elif suffix == ".xlsx":
        text = load_xlsx_bytes(data)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return LoadedDocument(name=name, text=text)


def combine_pasted_and_uploaded(pasted: str, uploads: list[Any] | None) -> list[LoadedDocument]:
    docs: list[LoadedDocument] = []
    if pasted and pasted.strip():
        docs.append(LoadedDocument(name="pasted_text", text=pasted.strip()))
    for uploaded in uploads or []:
        docs.append(load_uploaded_file(uploaded))
    return docs
